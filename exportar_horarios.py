
import os
import django
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

os.environ.setdefault("DJANGO_SETTINGS_MODULE", "citahn.settings")
django.setup()

from apps.medicos.models import Disponibilidad

libro = Workbook()
hoja = libro.active
hoja.title = "Horarios médicos"

hoja.append([
    "Identidad", "Nombre del médico", "Hospital",
    "Especialidad", "N.º colegiación", "Día",
    "Hora de inicio", "Hora de fin",
    "Horario activo", "Médico activo", "Usuario activo"
])

horarios = Disponibilidad.objects.select_related(
    "medico__usuario",
    "medico__hospital",
    "medico__especialidad"
).order_by(
    "medico__hospital__nombre",
    "medico__usuario__apellido",
    "dia_semana",
    "hora_inicio"
)

total = 0

for h in horarios.iterator():
    medico = h.medico
    usuario = medico.usuario

    hoja.append([
        usuario.identidad,
        f"{usuario.nombre} {usuario.apellido}",
        medico.hospital.nombre,
        medico.especialidad.nombre,
        medico.numero_colegiacion,
        h.get_dia_semana_display(),
        h.hora_inicio.strftime("%H:%M"),
        h.hora_fin.strftime("%H:%M"),
        "Sí" if h.activo else "No",
        "Sí" if medico.activo else "No",
        "Sí" if usuario.activo else "No",
    ])
    total += 1

for celda in hoja[1]:
    celda.font = Font(bold=True, color="FFFFFF")
    celda.fill = PatternFill("solid", fgColor="12345A")
    celda.alignment = Alignment(horizontal="center")

hoja.freeze_panes = "A2"
hoja.auto_filter.ref = hoja.dimensions

for columna in hoja.columns:
    letra = get_column_letter(columna[0].column)
    ancho = max(len(str(c.value or "")) for c in columna)
    hoja.column_dimensions[letra].width = min(ancho + 3, 35)

ruta = Path(__file__).resolve().parent / "planilla_medicos_citahn.xlsx"
libro.save(ruta)

print("Archivo:", ruta)
print("Horarios exportados:", total)
print("Filas de datos en la hoja:", hoja.max_row - 1)
print("Total de filas esperadas:", total)
